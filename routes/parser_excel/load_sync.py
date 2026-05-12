from .go_parse import go_parse
import os
import re
# Для работы в синхронном режиме (celery)


import xlrd,re
from pprint import pprint


import os
#import xlrd
from openpyxl import load_workbook
import csv
import asyncio

def print_error(error):
    return {
        'success':False,
        'errors':[error],
        'data':[]
    }


    #except Exception as e:
    #    print('Ошибка для при проверке ',s,type(s), str(e))
        #quit()


            



def parse_str(row):
    _str=[] ; hash_str={} ; col=0
    is_empty_cnt=0



def go_parse(**xarg):
    errors = []
    filename = xarg.get('filename')
    tmp_dir = xarg.get('tmp_dir')
    full_path = f"{tmp_dir}/{filename}"

    hash_fields = xarg.get('hash_fields')
    loopback = xarg.get('loopback')
    before_loopback = xarg.get('before_loopback')  # исправил опечатку: beore -> before
    data_line_number = xarg.get('data_line_number')
    limit = xarg.get('limit')

    # Определение расширения файла
    ext = os.path.splitext(filename)[1].lower()

    if ext not in ['.xls', '.xlsx', '.csv']:
        return print_error(f"Неподдерживаемое расширение файла: {ext}")

    # Чтение данных в виде строк (генератороподобно)
    try:
        if ext == '.xls':
            workbook = xlrd.open_workbook(full_path)
            sh = workbook.sheet_by_index(0)
            rows = (sh.row(row_idx) for row_idx in range(sh.nrows))
        elif ext == '.xlsx':
            wb = load_workbook(full_path, read_only=True)
            sh = wb.active
            rows = sh.iter_rows(values_only=True)
        elif ext == '.csv':
            def open_csv():
                encodings = ['utf-8', 'cp1251', 'latin1']
                for enc in encodings:
                    try:
                        f = open(full_path, 'r', encoding=enc, newline='')
                        sample = f.read(1024)
                        f.close()
                        f = open(full_path, 'r', encoding=enc, newline='')
                        return f, enc
                    except UnicodeDecodeError:
                        continue
                raise Exception("Не удалось определить кодировку CSV")
            
            file_obj, encoding = open_csv()
            # Указываем разделитель вручную
            reader = csv.reader(file_obj, delimiter=';')
            rows = reader
        else:
            return print_error("Логика не должна доходить до сюда")
    except Exception as e:
        return print_error(f"Ошибка при чтении файла {str(e)} (go_parse)")

    # Переменные состояния
    line_number = 0
    data = []
    cnt_empty_str = 0
    inserted_records = 0
    updated_records = 0

    try:
        for row in rows:
            # Пропуск строк до data_line_number
            if data_line_number and line_number < data_line_number:
                line_number += 1
                continue

            # Преобразуем строку в список значений
            if ext in ['.xls']:
                values = [str(cell.value) for cell in row]
            elif ext == '.xlsx':
                values = [str(cell) if cell is not None else '' for cell in row]
            elif ext == '.csv':
                values = [str(cell) if cell is not None else '' for cell in row]

            # Формируем hash_str или _str
            hash_str = {}
            _str = []

            for col, v in enumerate(values):
                #v=v.replace('№','&#8470;')
                if hash_fields:
                    if name := hash_fields.get(col):
                        hash_str[name] = v
                else:
                    _str.append(v)

            # Логика обработки
            if loopback:
                # Проверка на пустую строку
                if is_empty_xls_str(hash_str):
                    cnt_empty_str += 1
                else:
                    cnt_empty_str = 0

                # Вызов before_loopback, если задан
                if before_loopback:
                    before_loopback(hash_str)

                operation = loopback(hash_str)
                if operation == 'update':
                    updated_records += 1
                elif operation == 'insert':
                    inserted_records += 1

            else:
                _str = clean_empty_tail(_str)
                data.append(_str)

            # Прерывание при большом количестве пустых строк
            if cnt_empty_str > 20:
                print('is_empty str > 20')
                break

            line_number += 1
            if limit and line_number >= limit:
                print('limit: ', limit)
                break

    except Exception as e:
        return print_error(f"Ошибка при обработке строки: {str(e)}")
    finally:
        # Закрытие файла, если это CSV
        if ext == '.csv' and 'file_obj' in locals():
            file_obj.close()

    message = f"Добавлено записей: {inserted_records}<br>Обновлено записей: {updated_records}"

    return {
        'success': len(errors) == 0,
        'errors': errors,
        'loaded_filename': filename,
        'data': data,
        'message': message
    }

def print_error(error):
    return {
        'success':False,
        'errors':[error],
    }



