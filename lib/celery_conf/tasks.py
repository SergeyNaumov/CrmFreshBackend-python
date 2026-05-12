import json
import os
import re
from celery import shared_task
from datetime import datetime, timedelta
from decimal import Decimal
from pathlib import Path

import requests
from dotenv import load_dotenv
from jinja2 import Template
from openpyxl.utils import get_column_letter
from openpyxl.workbook import Workbook

from lib.celery_conf.config import celery_app, get_redis_client, BL_API_KEYS
from lib.celery_conf.utils import *

import sys
base_dir = Path(__file__).resolve().parent.parent.parent
sys.path.append(str(base_dir))

from db import get_db
from lib.send_mes import send_mes_task
from libsync.core_crm import get_owner
from config import config


from .parser_excel_tasks import parser_excel_load_sync


load_dotenv()

BR_TABLE = 'beeline_records_2'


def get_sub_from_redis(redis_client, sub_key):
    sub = redis_client.get(sub_key)
    return sub.decode("utf-8") if sub else None


@celery_app.task()
def check_xsi_sub(sub_key, api_key, URL):
    """проверяет статус подписки и обновляет, если надо"""
    redis_client = get_redis_client()
    sub = get_sub_from_redis(redis_client, sub_key)
    s = requests.Session()
    s.headers = {
        'X-MPBX-API-AUTH-TOKEN': api_key
    }
    resp = s.get(f'https://cloudpbx.beeline.ru/apis/portal/subscription?subscriptionId={sub}')
    if resp.status_code != 200 or not sub:
        data = {
            "expires": 90000,
            "subscriptionType": "BASIC_CALL",
            "url": URL
        }
        put_resp = s.put('https://cloudpbx.beeline.ru/apis/portal/subscription', json=data)
        if put_resp.status_code != 503:
            put_resp = put_resp.json()
            new_sub = put_resp.get('subscriptionId')
            if new_sub:
                redis_client.set(sub_key, new_sub)
            return new_sub
        return f'Подписка: {sub}. PUT cloudpbx.beeline.ru вернул код {put_resp.status_code}'
    return f'Подписка: {sub}. GET cloudpbx.beeline.ru вернул код {resp.status_code}'


@celery_app.task()
def renew_xsi_sub(sub_key, api_key, URL):
    """принудительно обновляет подписку раз в день"""
    s = requests.Session()
    s.headers = {
        'X-MPBX-API-AUTH-TOKEN': api_key
    }
    data = {
        "expires": 90000,
        "subscriptionType": "BASIC_CALL",
        "url": URL
    }
    put_resp = s.put('https://cloudpbx.beeline.ru/apis/portal/subscription', json=data)
    if put_resp.status_code == 503:
        return f'PUT cloudpbx.beeline.ru вернул код {put_resp.status_code}'
    put_resp = put_resp.json()
    new_sub = put_resp.get('subscriptionId')
    if new_sub:
        redis_client = get_redis_client()
        redis_client.set(sub_key, new_sub)
    return new_sub


@celery_app.task()
def xsi_data_handler(api_key_number, data):
    db = get_db(sync=1)
    event = data.get('xsi:Event')
    db.save(table='beeline_records_2_logs', data={'api_key_number': api_key_number, 'call_data': json.dumps(event)})
    event_data = event.get('xsi:eventData')
    call = event_data.get('xsi:call')
    if call is None:
        return 'None Call'
    call_id = call.get('xsi:callId')
    if call_id is None:
        return 'None CallId'
    ext_tracking_id = call.get('xsi:extTrackingId')
    call_state = call.get('xsi:state')
    is_sip, manager_raw_phone_or_sip = get_manager_phone_from_id(event.get('xsi:targetId'))
    beeline_abonent = get_beeline_abonent(db, manager_raw_phone_or_sip, is_sip)
    if beeline_abonent is None:
        return f"менеджер не найден. {event.get('xsi:targetId')}"
    try:
        client_phone = get_client_phone_from_remote_party(call.get('xsi:remoteParty'))
    except AttributeError:
        client_phone = None
    if call_state == 'Alerting' and client_phone is None:
        return "Wrong or None client phonenumber"
    direction = get_call_direction(call.get('xsi:personality'))
    call_data = {
        'externalId': ext_tracking_id,
        'beeline_abonent_id': beeline_abonent['id'],
        'api_key_number': api_key_number,
    }
    call_obj, call_phone, created, get_call_error = get_call_by_call_id(db, call_data)

    if call_obj is None:
        return get_call_error

    client_phone = call_phone if call_phone else client_phone
    data_to_save = {
        'callId': call_id, 'direction': direction, 'manager_id': beeline_abonent['manager_id']
    }
    if not is_sip:
        data_to_save['phone'] = client_phone
    db.save(
        table=BR_TABLE,
        data=data_to_save,
        update=1,
        where=f"id={call_obj}"
    )
    if created:
        if client_phone is None:
            return 'Wrong or None client phonenumber'
    else:
        if call_state == 'Alerting':
            return 'Call already exists'
        if call_state == 'Active':
            return 'Call already active or end'

    start_date, answer_date, release_date = None, None, None
    if call.get('xsi:startTime'):
        start_date = datetime.fromtimestamp(int(call.get('xsi:startTime')) / 1000)

    if call.get('xsi:answerTime'):
        answer_date = datetime.fromtimestamp(int(call.get('xsi:answerTime')) / 1000)

    if call.get('xsi:releaseTime'):
        release_date = datetime.fromtimestamp(int(call.get('xsi:releaseTime')) / 1000)

    date_data = {}
    if call_state == 'Alerting':
        date_data['start_date'] = start_date
    if call_state == 'Active':
        date_data['start_date'] = start_date
        date_data['answer_date'] = answer_date
    if call_state == 'Released':
        if answer_date:
            date_data['start_date'] = start_date
            date_data['answer_date'] = answer_date
            date_data['duration'] = int((release_date - answer_date).total_seconds())
        else:
            if direction == 'OUTBOUND':
                date_data['start_date'] = start_date
        date_data['release_date'] = release_date
    if date_data:
        db.save(
            table=BR_TABLE,
            update=1,
            data=date_data,
            where=f"id={call_obj}"
        )
    return "DONE"


@celery_app.task(time_limit=1560)
def search_records(api_keys):
    db = get_db(sync=1)
    start_date = datetime.today() - timedelta(days=6)
    calls = db.query(
        query=f"SELECT br.id, br.externalId, ba.phone, ba.sip_number, date(br.start_date) start_date FROM {BR_TABLE} br "
              f"LEFT JOIN beeline_abonent ba ON ba.id=br.beeline_abonent_id "
              f"WHERE downloaded=0 "
              f"AND start_date >= '{start_date.strftime('%Y-%m-%d %H:%M:%S')}' "
              f"AND duration > 0"
    )
    for api_key in api_keys:
        s = requests.Session()
        s.headers = {
            'X-MPBX-API-AUTH-TOKEN': api_key
        }
        _download_dir = './files/beeline'
        for call in calls:
            try:
                phone = call['sip_number'] if call['sip_number'] else call['phone'].replace('+7', '')
                response = s.get(f"https://cloudpbx.beeline.ru/apis/portal/v2/records/{call['externalId']}/{phone}/download")
            except Exception as e:
                continue
            if response.status_code == 200:
                date_dir = str(call['start_date']).replace('-', '/')
                download_dir = f"{_download_dir}/{date_dir}"

                full_name = f"{download_dir}/{call['id']}.mp3"
                os.makedirs(download_dir, exist_ok=True)

                with open(full_name, 'wb') as file:
                    file.write(response.content)

                    db.query(
                        query=f"UPDATE {BR_TABLE} set downloaded=1 where id=%s",
                        values=[call['id']]
                    )
    recursive_chmod('/var/www/fas-beeline/download')


@celery_app.task()
def sync_abonents_task():
    db = get_db(sync=1)
    db.query(query='delete ba from beeline_abonent ba left join manager m on m.id=ba.manager_id where m.gone=1')
    sync_abonent(db, [BL_API_KEYS[1], BL_API_KEYS[2], BL_API_KEYS[3]])


@celery_app.task()
def delete_old_records():
    db = get_db(sync=1)
    start_date = datetime.today() - timedelta(days=30)
    db.query(query=f"DELETE FROM {BR_TABLE} WHERE start_date <= '{start_date.strftime('%Y-%m-%d %H:%M:%S')}'")
    db.query(query=f"DELETE FROM beeline_records_2_logs WHERE registered <= '{start_date.strftime('%Y-%m-%d %H:%M:%S')}'")


@celery_app.task()
def delete_duplicate_users():
    # echo "from lib.celery_conf.tasks import delete_duplicate_users; delete_duplicate_users.delay()" | python
    db = get_db(sync=1)
    query = """
    SELECT
    u.id user_id, u.brand_id, MAX(um.registered) last_comment_date, u.contact_date, u.registered registered,
    u.inn, br2.id downloaded, (b.id or pi.inn) b_paid, uc.id AS has_contact
    FROM user u
    LEFT JOIN user_memo um ON um.user_id=u.id
    LEFT JOIN paided_inn pi ON u.inn=pi.inn
    LEFT JOIN user_contact uc ON uc.user_id=u.id and (uc.email<>'' or uc.phone<>'')
    LEFT JOIN beeline_records_2 br2 ON br2.phone=uc.phone and br2.duration>10
    LEFT JOIN docpack dp ON dp.user_id=u.id
    LEFT JOIN bill b ON b.docpack_id=dp.id and b.paid=1
    WHERE archive=0 AND u.inn <> ''
    GROUP BY u.id ORDER BY u.brand_id, u.inn
    """
    users_data = db.query(query=query)
    companies_by_brand_and_inn = {}
    for user in users_data:
        if len(user['inn'].strip()) not in [10, 12]:
            continue
        key = (user['brand_id'], user['inn'].strip())
        if key not in companies_by_brand_and_inn:
            companies_by_brand_and_inn[key] = []
        companies_by_brand_and_inn[key].append(user)

    today = datetime.now()
    min_contact_date = today - timedelta(days=365)
    max_contact_date = today + timedelta(days=365)

    to_archive = []
    for key, duplicates in companies_by_brand_and_inn.items():
        if len(duplicates) <= 1:
            continue

        duplicates.sort(
            key=lambda u: (
                u['b_paid'] is not None and u['b_paid'] == 1,
                u['downloaded'] or False,
                u['has_contact'] or False,
                min_contact_date <= parse_date(u['contact_date']) <= max_contact_date if u['contact_date'] else False,
                parse_date(u['last_comment_date']),
            ),
            reverse=True
        )
        to_archive.extend(
            u for u in duplicates[1:]
            if not (
                u['b_paid'] is not None and u['b_paid'] == 1
                or (u['downloaded'] is not None)
                or (u['has_contact'] is not None)
                or (u['contact_date'] and min_contact_date <= parse_date(u['contact_date']) <= max_contact_date)
            )
        )

    to_archive_ids = [u['user_id'] for u in to_archive]
    if to_archive_ids:
        values = ", ".join(f"({u})" for u in to_archive_ids)
        db.query(query=f"""INSERT INTO log_move_user_to_archive (id) VALUES {values}""")
        db.save(
            table='user',
            data={"archive": 1},
            update=1,
            where=f"id IN ({','.join(map(str, to_archive_ids))})"
        )
    return f'В архив отправлено: {len(to_archive_ids)}'


@celery_app.task()
def user_contact_clean():
    db = get_db(sync=1)
    existing_phones_data = db.query(
        query='SELECT phone, user_id FROM user_contact'
    )

    existing_phones_dict = {}
    for row in existing_phones_data:
        existing_phones_dict.setdefault(row['user_id'], set()).add(row['phone'])

    all_values = []
    data = db.query(query='select id,user_id,phone from user_contact where length(phone)>12')
    for row in data:
        phones = format_phones(row['phone'])
        if phones:
            user_id = row['user_id']
            existing_phones_set = existing_phones_dict.get(user_id, set())

            phones_to_create = [
                {"phone": p, "user_id": user_id}
                for p in phones[1:]
                if p not in existing_phones_set
            ]
            all_values.extend(phones_to_create)

            db.save(
                table='user_contact', update=1, data={"phone": phones[0]}, where=f"id={row['id']}"
            )
            existing_phones_dict[user_id].add(phones[0])

    if all_values:
        values = ", ".join(
            f"({entry['user_id']}, '{entry['phone']}')"
            for entry in all_values
        )
        db.query(query=f'''INSERT INTO user_contact (user_id, phone) VALUES {values}''')


@celery_app.task()
def clear_mail_send_table():
    db = get_db(sync=1)
    five_days_ago = datetime.today() - timedelta(days=5)
    db.query(query=f"DELETE FROM mail_send WHERE registered < '{five_days_ago}'")
    db.query(query=f"optimize table mail_send")


@celery_app.task()
def check_user_input_work():
    # echo "from lib.celery_conf.tasks import check_user_input_work; check_user_input_work.delay()" | python
    db = get_db(sync=1)
    calls = db.query(
        query=f"""
        SELECT br2.duration, br2.phone, br2.start_date, u.id user_id  
        FROM beeline_records_2 br2 
        LEFT JOIN beeline_abonent ba ON ba.id=br2.beeline_abonent_id 
        LEFT JOIN user_contact uc ON uc.phone=br2.phone
        LEFT JOIN user u ON uc.user_id=u.id AND u.manager_id=ba.manager_id
        WHERE u.id IS NOT NULL AND br2.duration > 120 AND br2.phone <> '' and u.id NOT IN (SELECT user_id FROM user_input_work)
        """
    )
    users = {}
    for call in calls:
        user_id = call['user_id']
        phone = call['phone']
        duration = call['duration']
        start_date = call['start_date']

        if user_id not in users:
            users[user_id] = {}

        if phone not in users[user_id]:
            users[user_id][phone] = {"duration": duration, "start_date": start_date}
        else:
            if duration > users[user_id][phone]["duration"]:
                users[user_id][phone] = {"duration": duration, "start_date": start_date}

    results = {}
    for card_id, phones in users.items():
        max_call = max(phones.values(), key=lambda x: x["duration"])
        results[card_id] = max_call["start_date"]

    if results:
        values = ", ".join(
            f"({key}, '{value}')"
            for key, value in results.items()
        )
        db.query(query=f"""INSERT INTO user_input_work (user_id, ts) VALUES {values}""")
    return f'Введено в работу: {len(results)}'


@celery_app.task()
def generate_xlsx_with_unused_phones():
    db = get_db(sync=1)
    current_date = datetime.now()
    next_day = current_date + timedelta(days=1)
    if next_day.day != 1:
        return 'not today'
    manager_data = db.query(
        query=f"""
            select m.id manager_id, m.name, m.phone manager_phone, m.registered, MAX(s.registered) last_login, br2.id has_records
            from manager m 
            left join beeline_abonent ba on ba.manager_id=m.id 
            left join beeline_records_2 br2 on br2.beeline_abonent_id=ba.id and br2.start_date >= '{current_date.strftime('%Y-%m-01')}'
            left join session s on s.auth_id=m.id
            where m.gone=0 and m.phone <> ''
            group by m.id
            having has_records is NULL
        """
    )
    wb = Workbook()
    ws = wb.active

    headers = ['id', 'Имя менеджера', 'Номер телефона', 'Зарегистрирован', 'Последний вход в систему']
    for col_num, header in enumerate(headers, 1):
        cell = ws.cell(row=1, column=col_num)
        cell.value = header
    column_widths = [15, 50, 20, 30, 30]
    for col_num, width in enumerate(column_widths, 1):
        col_letter = get_column_letter(col_num)
        ws.column_dimensions[col_letter].width = width

    for row_num, row_data in enumerate(manager_data, 2):
        ws.cell(row=row_num, column=1, value=row_data['manager_id'])
        ws.cell(row=row_num, column=2, value=row_data['name'])
        ws.cell(row=row_num, column=3, value=row_data['manager_phone'])
        ws.cell(row=row_num, column=4, value=row_data['registered'].strftime("%Y-%m-%d %H:%M:%S") if row_data['registered'] else '')
        ws.cell(row=row_num, column=5, value=row_data['last_login'].strftime("%Y-%m-%d %H:%M:%S") if row_data['last_login'] else '')

    tmp_dir = 'tmp'
    os.makedirs(tmp_dir, exist_ok=True)
    output_filename = f"{tmp_dir}/managers_{current_date.strftime('%Y_%m_%d')}.xlsx"
    wb.save(output_filename)
    try:
        send_mes_task({
            'from_addr': 'no-reply@fascrm.ru',
            'to': 'zer@t-pass.pro,ship@t-pass.pro,pzm@t-pass.pro,dimitri.k@t-pass.pro',
            'subject': f"Уведомление о неиспользуемых телефонных номерах",
            'message': 'См. вложение',
            'attachment': output_filename,
        })
    except Exception as e:
        print(f'generate_xlsx_with_unused_phones send email error: {str(e)}')
    os.remove(output_filename)


@celery_app.task()
def check_phones_in_identifiers():
    api_phone_url = config['API_PHONE_URL']
    headers = {
        'token': config['API_PHONE_TOKEN']
    }
    source_id = 2  # fas
    limit, offset = 1000, 0
    all_data = []
    while True:
        r = requests.get(url=f"{api_phone_url}/each?limit={limit}&offset={offset}&source_id={source_id}", headers=headers, verify=False)
        if r.status_code != 200:
            return f'{r.status_code=}'
        data = r.json()
        all_data.extend(data)
        if len(data) < 1000:
            break
        else:
            limit += 1000
            offset += 1000
    all_phones_dict = {obj['number']: obj['id'] for obj in all_data}
    all_phones = set(all_phones_dict.keys())
    db = get_db(sync=1)
    manager_phones = db.query(query=f"select phone, gone from manager where phone <> ''")

    not_gone_phones = []
    gone_phones = []
    for obj in manager_phones:
        obj['phone'] = parse_phone(obj['phone']).replace('+', '')
        if not obj['phone'].startswith('7'):
            obj['phone'] = f"7{obj['phone']}"
        if obj['gone'] == 0:
            not_gone_phones.append(obj['phone'])
        else:
            gone_phones.append(obj['phone'])

    not_gone_phones = set(not_gone_phones)
    gone_phones = set(gone_phones)

    missing_phones = list(not_gone_phones - all_phones)
    for phone in missing_phones:
        r = requests.post(url=f"{api_phone_url}/create", headers=headers, data=json.dumps({"nums": phone, "source_id": source_id}), verify=False)

    # need_to_delete_phones = list(gone_phones & all_phones)
    # for phone in need_to_delete_phones:
    #     phone_id = all_phones_dict.get(phone)
    #     if phone_id:
    #         r = requests.delete(url=f"{api_phone_url}/delete/{phone_id}", headers=headers, verify=False)


@celery_app.task()
def teamwork_ofp_wins_send():
    db = get_db(sync=1)
    data_to_send = db.query(query="""
        select u.firm, u.inn, m.name, m2.name lawyer, top.header product_name, b.header brand_name,
        group_concat(concat(uc.fio, '№',uc.email,'№', uc.phone) SEPARATOR '~') contact_info, ofp.teamwork_ofp_id
        from teamwork_ofp ofp
        left join user u on u.id=ofp.user_id
        left join brand b on b.id=u.brand_id
        join user_contact uc on uc.user_id=u.id and (uc.phone <> '' or uc.email <> '')
        left join manager m on m.id=ofp.manager_from
        left join manager m2 on m2.id=ofp.manager_to
        left join teamwork_ofp_wins tow on tow.teamwork_ofp_id=ofp.teamwork_ofp_id
        left join teamwork_ofp_product top on top.id=ofp.product
        where win_status=1 and (tow.sent=0 or tow.sent is null) and u.id is not null
        group by ofp.teamwork_ofp_id
    """)
    for obj in data_to_send:
        obj['contact_data'] = ''
        for info in obj['contact_info'].split('~'):
            if info:
                fio, email, phone = info.split('№')
                fio = fio if fio else 'Не указано'
                email = email if email else 'Не указано'
                phone = phone if phone else 'Не указано'
                if obj['contact_data']:
                    obj['contact_data'] += '<br>'
                obj['contact_data'] += f'Контакт: {fio},<br> Email: {email},<br> Телефон: {phone}<br>'
        for key, value in obj.items():
            if value is None:
                obj[key] = 'Не указано'

    if data_to_send:
        with open('lib/celery_conf/template/teamwork_ofp_wins.html', 'r', encoding='utf-8') as file:
            template = file.read()
        t = Template(template)
        result = t.render(data_to_send=data_to_send)
        send_mes_task.delay({
            'from_addr': 'no-reply@fascrm.ru',
            'to': 'tech@t-pass.pro',
            'subject': "Победы за сегодня",
            'message': result,
        })
        ids = [obj['teamwork_ofp_id'] for obj in data_to_send]
        values = ", ".join(f"({t}, 1)" for t in ids)
        db.query(query=f"""insert into teamwork_ofp_wins (teamwork_ofp_id, sent) values {values}""")


@celery_app.task()
def teamwork_ofp_week_wins_send():
    db = get_db(sync=1)
    week_ago = datetime.today() - timedelta(days=7)
    data_to_send = db.query(query=f"""
        select u.firm, u.inn, m.name, m2.name lawyer, top.header product_name, b.header brand_name,
        group_concat(concat(uc.fio, '№',uc.email,'№', uc.phone) SEPARATOR '~') contact_info, ofp.teamwork_ofp_id
        from teamwork_ofp ofp
        left join user u on u.id=ofp.user_id
        left join brand b on b.id=u.brand_id
        join user_contact uc on uc.user_id=u.id and (uc.phone <> '' or uc.email <> '')
        left join manager m on m.id=ofp.manager_from
        left join manager m2 on m2.id=ofp.manager_to
        left join teamwork_ofp_wins tow on tow.teamwork_ofp_id=ofp.teamwork_ofp_id
        left join teamwork_ofp_product top on top.id=ofp.product
        where win_status=1 and tow.sent=1 and tow.created_at >= '{week_ago.strftime("%Y-%m-%d 00:00:00")}' and u.id is not null
        group by ofp.teamwork_ofp_id
    """)
    for obj in data_to_send:
        obj['contact_data'] = ''
        for info in obj['contact_info'].split('~'):
            if info:
                fio, email, phone = info.split('№')
                fio = fio if fio else 'Не указано'
                email = email if email else 'Не указано'
                phone = phone if phone else 'Не указано'
                if obj['contact_data']:
                    obj['contact_data'] += '<br>'
                obj['contact_data'] += f'Контакт: {fio},<br> Email: {email},<br> Телефон: {phone}<br>'
        for key, value in obj.items():
            if value is None:
                obj[key] = 'Не указано'

    if data_to_send:
        with open('lib/celery_conf/template/teamwork_ofp_wins.html', 'r', encoding='utf-8') as file:
            template = file.read()
        t = Template(template)
        result = t.render(data_to_send=data_to_send)
        send_mes_task.delay({
            'from_addr': 'no-reply@fascrm.ru',
            'to': 'davydov_igor@t-pass.pro',
            'subject': "Победы за прошлую неделю",
            'message': result,
        })


@celery_app.task()
def send_current_kpi_traffic_notifications():
    db = get_db(sync=1)
    today = datetime.now().date()
    if day_off := db.query(query="select date from day_off where date=%s", onevalue=1, values=[today]):
        return f"{day_off} - выходной"
    managers = db.query(
        query=f"""
        SELECT m.id manager_id, name, kpi_traffic FROM manager m
        left join manager_group mg on mg.id=m.group_id
        where mg.show_kpi=1 and m.gone=0 and kpi_traffic <> 0;
        """
    )
    if not managers:
        return 'Нет менеджеров'
    manager_data = {manager['manager_id']: manager for manager in managers}

    calls = db.query(
        query=f"""
        SELECT sum(br2.duration) duration, br2.manager_id
        FROM beeline_records_2 br2
        WHERE br2.duration <> 0 and br2.start_date >= %s and br2.start_date <= %s and br2.manager_id in ({','.join(map(str, manager_data.keys()))})
        group by br2.manager_id
        """,
        values=[f"{today} 00:00:00", f"{today} 23:59:59"]
    )
    call_data = {call['manager_id']: call['duration'] for call in calls}
    for m_id in manager_data.keys():
        if m_id not in call_data:
            call_data[m_id] = Decimal(0)
    owner_data = {}
    subject = "Трафик KPI Fas"
    for m_id, m_data in manager_data.items():
        name = m_data['name']
        manager_email = get_email(db, m_id)
        today_traffic_p = round(100 * call_data[m_id] / (m_data['kpi_traffic'] * 60), 2)
        message = (f"{name}, Ваш трафик {format_sec(call_data[m_id])}, что составляет "
                   f"{today_traffic_p}% от нормативного дневного трафика в {m_data['kpi_traffic']} мин.<br>")
        owner = get_owner(db=db, manager_id=m_id)
        owner_email = owner.get('email')
        if not owner_email:
            owner_email = get_email(db, owner['id'])

        if owner['id'] != m_id:
            send_mes_task.delay({
                'from_addr': 'no-reply@fascrm.ru',
                'to': f'{manager_email}',
                'subject': subject,
                'message': message,
            })

        if owner_email:
            if owner_data.get(owner_email):
                owner_data[owner_email] += message
            else:
                owner_data[owner_email] = message

    for own_email, messages in owner_data.items():
        send_mes_task.delay({
            'from_addr': 'no-reply@fascrm.ru',
            'to': f'{own_email},zer@t-pass.pro,ship@t-pass.pro,pzm@t-pass.pro',
            'subject': subject,
            'message': messages,
        })


@celery_app.task()
def add_info_to_sip_calls():
    one_day_ago = datetime.now() - timedelta(days=1)
    db = get_db(sync=1)
    exists_calls = db.query(
        query=f"""
        select br2.id from {BR_TABLE} br2 left join beeline_abonent ba on ba.id=br2.beeline_abonent_id 
        where ba.sip_number <> '' and br2.phone = '' and start_date >= '{one_day_ago.strftime('%Y-%m-%d %H:%M:%S')}' and br2.duration > 0
    """, onevalue=1
    )
    if exists_calls:
        abonents = db.query(
            query=f"""select ba.manager_id id, ba.sip_number, ba.id beeline_abonent_id from beeline_abonent ba where sip_number <> ''"""
        )
        if abonents:
            add_info_to_sip_calls_from_date(db, abonents, [BL_API_KEYS[1], BL_API_KEYS[2], BL_API_KEYS[3]], 24)

#from .parser_excel_tasks import parser_excel_go_parse
